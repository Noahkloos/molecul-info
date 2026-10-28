"""Word-document (.docx) van de Hazardous Substance Risk Assessment.

De opmaak volgt het schoolsjabloon: een titelregel met Ref./titel/stofnaam, blauwe
sectiebalken, de 9 GHS-pictogrammen met een X bij de pictogrammen die gelden, en
aankruisvakjes voor gevarentype, blootstellingsroute en PBM. Het document is bedoeld om
in een labverslag te plakken (alles staat in tabellen, dus het blijft netjes bij kopiëren).
"""

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .risico import ERNST_LABEL, GEVARENTYPEN, KANS_LABEL, PERSONEN, PICTOGRAM_NAMEN, ROUTES, Risicobeoordeling

PICTO_MAP = Path(__file__).resolve().parent.parent / "data" / "ghs_pictogrammen"
BALK_KLEUR = "BDD7EE"  # lichtblauwe sectiebalk zoals in het sjabloon
AAN, UIT = "☒", "☐"


# --------------------------------------------------------------------------- #
# Opmaakhulpjes (python-docx heeft geen directe functie voor celkleur en randen)
# --------------------------------------------------------------------------- #
def _celkleur(cel, kleur: str) -> None:
    tcPr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kleur)
    tcPr.append(shd)


def _randen(tabel) -> None:
    tblPr = tabel._tbl.tblPr
    randen = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right", "insideH", "insideV"):
        rand = OxmlElement(f"w:{kant}")
        rand.set(qn("w:val"), "single")
        rand.set(qn("w:sz"), "6")
        rand.set(qn("w:color"), "000000")
        randen.append(rand)
    tblPr.append(randen)


def _tekst(cel, tekst: str, vet: bool = False, cursief: bool = False, grootte: float | None = None,
           uitlijning=None, nieuw: bool = False):
    alinea = cel.add_paragraph() if nieuw else cel.paragraphs[0]
    run = alinea.add_run(tekst)
    run.bold, run.italic = vet, cursief
    if grootte:
        run.font.size = Pt(grootte)
    if uitlijning is not None:
        alinea.alignment = uitlijning
    alinea.paragraph_format.space_after = Pt(1)
    return run


def _tabel(doc, rijen: int, kolommen: int, breedtes_cm: list[float]):
    """Tabel met vaste breedte (anders past Word de kolommen zelf aan en lopen de blokken niet gelijk)."""
    tabel = doc.add_table(rows=rijen, cols=kolommen)
    tabel.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabel.autofit = False
    _randen(tabel)
    tblPr = tabel._tbl.tblPr
    indeling = OxmlElement("w:tblLayout")
    indeling.set(qn("w:type"), "fixed")
    tblPr.append(indeling)
    breedte = OxmlElement("w:tblW")
    breedte.set(qn("w:w"), str(int(sum(breedtes_cm) * 567)))  # 1 cm = 567 twips
    breedte.set(qn("w:type"), "dxa")
    tblPr.append(breedte)
    for kolom, cm in zip(tabel.columns, breedtes_cm):
        kolom.width = Cm(cm)
    for rij in tabel.rows:
        for cel, cm in zip(rij.cells, breedtes_cm):
            cel.width = Cm(cm)
    return tabel


def _sectie(doc, titel: str, toelichting: str = ""):
    """Blauwe balk met titel (en cursieve toelichting), zoals 'Classification (state the category of danger)'."""
    tabel = _tabel(doc, 1, 1, [18.0])
    cel = tabel.cell(0, 0)
    _celkleur(cel, BALK_KLEUR)
    _tekst(cel, titel, vet=True)
    if toelichting:
        _tekst_run = cel.paragraphs[0].add_run(f" {toelichting}")
        _tekst_run.bold, _tekst_run.italic = True, True
    return tabel


def _vakjes(opties: list[str], aangevinkt: list[str]) -> str:
    return "     ".join(f"{AAN if o in aangevinkt else UIT} {o}" for o in opties)


def _blok(doc, regels: list[str] | str, leeg: str = "Not available in PubChem; consult the supplier SDS."):
    tabel = _tabel(doc, 1, 1, [18.0])
    cel = tabel.cell(0, 0)
    regels = [regels] if isinstance(regels, str) else regels
    regels = [r for r in regels if r and r.strip()] or [leeg]
    for i, regel in enumerate(regels):
        _tekst(cel, regel, nieuw=i > 0)
    return tabel


def _ruimte(doc) -> None:
    """Kleine witregel: voorkomt dat Word aangrenzende tabellen samenvoegt."""
    alinea = doc.add_paragraph()
    alinea.paragraph_format.space_after = Pt(0)
    alinea.paragraph_format.space_before = Pt(0)
    alinea.add_run().font.size = Pt(2)


# --------------------------------------------------------------------------- #
# Document
# --------------------------------------------------------------------------- #
def maak_docx(b: Risicobeoordeling) -> bytes:
    doc = Document()
    for sectie in doc.sections:
        sectie.left_margin = sectie.right_margin = Cm(1.5)
        sectie.top_margin = sectie.bottom_margin = Cm(1.5)
    stijl = doc.styles["Normal"]
    stijl.font.name = "Calibri"
    stijl.font.size = Pt(9.5)

    # Titelregel
    titel = _tabel(doc, 1, 3, [3.5, 9.5, 5.0])
    _tekst(titel.cell(0, 0), f"Ref. {b.ref}", vet=True, uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
    _tekst(titel.cell(0, 1), "Hazardous Substance Risk Assessment", vet=True, grootte=14,
           uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
    _tekst(titel.cell(0, 2), "Chemical Name", vet=True, uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
    _tekst(titel.cell(0, 2), b.chemische_naam, uitlijning=WD_ALIGN_PARAGRAPH.CENTER, nieuw=True)

    # Algemene gegevens
    algemeen = [
        ("Substance(s) & Manufacturer/s", b.stof_en_fabrikant),
        ("Location/s", b.locatie),
        ("Persons at Risk", _vakjes(PERSONEN, b.personen)),
        ("Task/s & Description", b.taak),
        ("Quantity used", b.hoeveelheid),
    ]
    info = _tabel(doc, len(algemeen), 2, [4.5, 13.5])
    for i, (label, waarde) in enumerate(algemeen):
        _tekst(info.cell(i, 0), label, vet=True)
        _tekst(info.cell(i, 1), waarde or "")
    _ruimte(doc)

    # Classificatie + pictogrammen
    _sectie(doc, "Classification", "(state the category of danger)")
    kop = _tabel(doc, 1, 1, [18.0])
    _tekst(kop.cell(0, 0), "Hazard pictograms – select all that apply to the work activity.", vet=True)
    pictos = _tabel(doc, 3, 9, [2.0] * 9)
    for i, (code, naam) in enumerate(PICTOGRAM_NAMEN.items()):
        cel = pictos.cell(0, i)
        cel.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        pad = PICTO_MAP / f"{code}.png"
        if pad.exists():
            cel.paragraphs[0].add_run().add_picture(str(pad), width=Cm(1.5))
        _tekst(pictos.cell(1, i), naam, grootte=7, uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
        _tekst(pictos.cell(2, i), "X" if code in b.pictogrammen else "", vet=True,
               uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
    signaal = _tabel(doc, 1, 2, [4.5, 13.5])
    _tekst(signaal.cell(0, 0), "Signal word", vet=True)
    _tekst(signaal.cell(0, 1), (b.signaalwoord or "None").upper(), vet=True)
    _ruimte(doc)

    # Gevarentype
    _sectie(doc, "Hazard Type")
    andere = [t for t in b.gevarentypen if t not in GEVARENTYPEN]
    _blok(doc, _vakjes(GEVARENTYPEN, b.gevarentypen) + f"     {AAN if andere else UIT} Other: {', '.join(andere)}")
    _ruimte(doc)

    # H- en P-zinnen
    _sectie(doc, "Hazard Statements")
    _blok(doc, [f"{c} – {t}" for c, t in b.h_zinnen] + ([f"Classification source: {b.ghs_bron}"] if b.ghs_bron else []),
          leeg="No GHS hazard statements found in PubChem. This does NOT mean the substance is safe; consult the SDS.")
    _ruimte(doc)
    _sectie(doc, "Precautionary Statements")
    _blok(doc, [f"{c} – {t}" for c, t in b.p_zinnen], leeg="No precautionary statements found; consult the SDS.")
    _ruimte(doc)

    # Blootstellingsroute
    _sectie(doc, "Main Route of Exposure")
    _blok(doc, _vakjes(ROUTES, b.routes))
    _ruimte(doc)

    # Grenswaarden
    _sectie(doc, "Control Parameters:", "Exposure Limits and Monitoring")
    _blok(doc, [f"• {g}" for g in b.grenswaarden] + [f"Monitoring: {b.monitoring}"])
    _ruimte(doc)

    # Beheersmaatregelen
    _sectie(doc, "Control Measures:", "Note engineering, ventilation and other requirements needed")
    maatregelen = _tabel(doc, 1, 1, [18.0])
    cel = maatregelen.cell(0, 0)
    _tekst(cel, "Engineering Measures", vet=True)
    for regel in b.technische_maatregelen:
        _tekst(cel, f"• {regel}", nieuw=True)
    _tekst(cel, "Hygiene Measures", vet=True, nieuw=True)
    for regel in b.hygiene:
        _tekst(cel, f"• {regel}", nieuw=True)
    _ruimte(doc)

    # PBM
    _sectie(doc, "Personal Protective Equipment", "(state type and standard)")
    pbm = _tabel(doc, len(b.pbm) + 1, 4, [1.2, 5.3, 4.5, 7.0])
    for i, kopje in enumerate(["", "Type", "Standard", "Notes"]):
        _tekst(pbm.cell(0, i), kopje, vet=True)
    for r, item in enumerate(b.pbm, start=1):
        _tekst(pbm.cell(r, 0), AAN if item.nodig else UIT, grootte=12, uitlijning=WD_ALIGN_PARAGRAPH.CENTER)
        _tekst(pbm.cell(r, 1), item.item, vet=item.nodig)
        _tekst(pbm.cell(r, 2), item.standaard)
        _tekst(pbm.cell(r, 3), item.toelichting)
    _ruimte(doc)

    # EHBO
    _sectie(doc, "First Aid Measures")
    ehbo = _tabel(doc, len(ROUTES), 2, [3.0, 15.0])
    for i, route in enumerate(ROUTES):
        _tekst(ehbo.cell(i, 0), route, vet=True)
        _tekst(ehbo.cell(i, 1), b.ehbo.get(route) or "Consult the supplier SDS.")
    _ruimte(doc)

    for titel_, inhoud in (("Fire Fighting", b.brand), ("Storage and Incompatibilities", b.opslag),
                           ("Spillage / Accidental Release", b.morsen), ("Waste Disposal", b.afval)):
        _sectie(doc, titel_)
        _blok(doc, inhoud)
        _ruimte(doc)

    # Fysische gegevens
    _sectie(doc, "Physical and Chemical Properties", "(for reference)")
    fys = _tabel(doc, len(b.fysisch), 2, [4.5, 13.5])
    for i, (label, waarde) in enumerate(b.fysisch):
        _tekst(fys.cell(i, 0), label, vet=True)
        _tekst(fys.cell(i, 1), waarde)
    _ruimte(doc)

    # Risicobeoordeling
    _sectie(doc, "Risk Rating", "(with the control measures above in place)")
    risico = _tabel(doc, 2, 3, [6.0, 6.0, 6.0])
    for i, kopje in enumerate(["Severity (1–3)", "Likelihood (1–3)", "Risk = Severity × Likelihood"]):
        _tekst(risico.cell(0, i), kopje, vet=True)
    _tekst(risico.cell(1, 0), f"{b.ernst} – {ERNST_LABEL[b.ernst]}")
    _tekst(risico.cell(1, 1), f"{b.kans} – {KANS_LABEL[b.kans]}")
    _tekst(risico.cell(1, 2), f"{b.risicoscore} – {b.risiconiveau.upper()}", vet=True)
    _ruimte(doc)

    # Ondertekening
    _sectie(doc, "Assessment")
    onder = _tabel(doc, 2, 4, [4.5, 4.5, 4.5, 4.5])
    for i, (kopje, waarde) in enumerate([("Assessed by", b.beoordelaar), ("Date", b.datum),
                                         ("Review date", b.herzieningsdatum), ("Signature", "")]):
        _tekst(onder.cell(0, i), kopje, vet=True)
        _tekst(onder.cell(1, i), waarde)
    _ruimte(doc)

    # Bronnen
    _sectie(doc, "Sources")
    bron_tabel = _tabel(doc, 1, 1, [18.0])
    cel = bron_tabel.cell(0, 0)
    for i, (naam, url) in enumerate(b.bronnen):
        run = _tekst(cel, f"• {naam}" + (f" – {url}" if url else ""), grootte=8, nieuw=i > 0)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    _tekst(cel, "Generated with Molecuul-uitlegger from public PubChem data. Always verify against the current "
                "supplier Safety Data Sheet (SDS) before starting the experiment.",
           cursief=True, grootte=8, nieuw=True)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
